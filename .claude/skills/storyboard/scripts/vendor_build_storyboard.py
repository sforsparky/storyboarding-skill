#!/usr/bin/env python3
"""
Build a VSL storyboard .xlsx from a JSON spec.

Usage:
    python3 build_storyboard.py input.json output.xlsx

Input JSON schema:
{
  "tabs": [
    {
      "tab_name": "Main",                 # sheet tab name (keep short & descriptive)
      "script_url": "",                   # optional, leave "" for placeholder text
      "source_folder_url": "",            # optional, leave "" for placeholder text
      "rows": [
        {
          "visual_type": "TALKING HEAD",  # must be one of VISUAL_TYPES keys below
          "script": "Line of dialogue\nSecond short line",
          "visual_direction": "",
          "notes": ""
        },
        ...
      ]
    },
    ...
  ]
}

Multiple tabs share ONE legend sheet, added at the end of the workbook.
"""

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Visual type -> color map. Keys must match what the caller puts in
# row["visual_type"] (case-insensitive match is applied).
# ---------------------------------------------------------------------------
VISUAL_TYPES = {
    "TALKING HEAD":               "2E75B6",  # blue
    "B-ROLL":                     "548235",  # green
    "CUSTOM GRAPHIC":             "7030A0",  # purple
    "SCREENCAST":                 "ED7D31",  # orange
    "TESTIMONIAL":                "D6408C",  # pink
    "TALKING HEAD + LOWER THIRD": "00A19A",  # teal
    "B-ROLL / GRAPHIC":           "8B5A2B",  # brown
    "GRAPHIC / SCREENCAST":       "5B7A99",  # blue-grey
}

TALKING_HEAD_TYPES = {"TALKING HEAD", "TALKING HEAD + LOWER THIRD"}

HEADER_FILL_COLORS = {
    "A": "1F2937",  # Line #    - near-black slate
    "B": "1F2937",  # Visual Placeholder
    "C": "1F2937",  # Script
    "D": "1F2937",  # Visual Direction
    "E": "1F2937",  # Notes
}
# Distinct-per-column header look: same dark base, but we vary a hint of hue
# per column so columns are still visually distinguishable at a glance.
HEADER_FILL_COLORS = {
    "A": "31363F",
    "B": "263A52",
    "C": "2E2E2E",
    "D": "3B2F4A",
    "E": "31363F",
}

ZEBRA_FILL = "F2F2F2"
COLUMN_WIDTHS = {"A": 8, "B": 30, "C": 46, "D": 40, "E": 30}
ROW_HEIGHT_DATA = 90
ROW_HEIGHT_HEADER = 22
ROW_HEIGHT_UTILITY = 24

WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def normalize_type(vtype: str) -> str:
    key = (vtype or "").strip().upper()
    if key in VISUAL_TYPES:
        return key
    # forgiving fallback for near-miss labels
    for k in VISUAL_TYPES:
        if k.replace(" ", "") == key.replace(" ", ""):
            return k
    raise ValueError(f"Unknown visual_type: {vtype!r}. Must be one of {list(VISUAL_TYPES)}")


def set_cell(ws, row, col_letter, value, *, font=None, fill=None, align=WRAP):
    cell = ws[f"{col_letter}{row}"]
    cell.value = value
    if align:
        cell.alignment = align
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    return cell


def build_utility_bar(ws, script_url, source_folder_url):
    row = 1
    ws.merge_cells(f"A{row}:B{row}")
    set_cell(ws, row, "A", "\U0001F4C1 PROJECT LINKS",
             font=Font(bold=True, color="FFFFFF"),
             fill=PatternFill("solid", fgColor="404040"),
             align=WRAP_CENTER)

    script_label = script_url if script_url else "\U0001F4C4 Original Script — paste URL here"
    c = set_cell(ws, row, "C", script_label,
                 font=Font(color="0563C1", underline="single"),
                 fill=PatternFill("solid", fgColor="FFFFFF"),
                 align=WRAP)
    if script_url:
        c.hyperlink = script_url

    ws.merge_cells(f"D{row}:E{row}")
    folder_label = source_folder_url if source_folder_url else "\U0001F4C2 Source Files Folder — paste URL here"
    d = set_cell(ws, row, "D", folder_label,
                 font=Font(color="1E7145", underline="single"),
                 fill=PatternFill("solid", fgColor="FFFFFF"),
                 align=WRAP)
    if source_folder_url:
        d.hyperlink = source_folder_url

    ws.row_dimensions[row].height = ROW_HEIGHT_UTILITY


def build_header(ws):
    row = 2
    headers = {
        "A": "Line #",
        "B": "Visual Placeholder",
        "C": "Script",
        "D": "Visual Direction",
        "E": "Notes",
    }
    for col, label in headers.items():
        set_cell(ws, row, col, label,
                 font=Font(bold=True, color="FFFFFF"),
                 fill=PatternFill("solid", fgColor=HEADER_FILL_COLORS[col]),
                 align=WRAP_CENTER)
    ws.row_dimensions[row].height = ROW_HEIGHT_HEADER


def build_rows(ws, rows):
    start_row = 3
    for i, row_data in enumerate(rows):
        r = start_row + i
        vtype = normalize_type(row_data["visual_type"])
        color = VISUAL_TYPES[vtype]
        is_talking_head = vtype in TALKING_HEAD_TYPES
        zebra = (i % 2 == 1)

        base_fill = PatternFill("solid", fgColor=ZEBRA_FILL) if zebra else None

        # Line #
        set_cell(ws, r, "A", i + 1, align=WRAP_CENTER, fill=base_fill)

        # Visual Placeholder (Column B) - always colored by visual type,
        # zebra striping never overrides this.
        set_cell(ws, r, "B", vtype,
                 font=Font(bold=True, color="FFFFFF"),
                 fill=PatternFill("solid", fgColor=color),
                 align=WRAP_CENTER)

        # Script
        set_cell(ws, r, "C", row_data.get("script", ""), fill=base_fill)

        # Visual Direction - text color matches visual type color. Talking-head
        # rows are blank except the single allowed cut-back note; any other
        # text supplied for a talking-head row is intentionally discarded.
        if is_talking_head:
            vdir = "Cut back to talking head." if row_data.get("cut_back") else ""
        else:
            vdir = row_data.get("visual_direction", "") or ""
        set_cell(ws, r, "D", vdir, font=Font(color=color), fill=base_fill)

        # Notes - production cues only
        notes = "" if is_talking_head else row_data.get("notes", "")
        set_cell(ws, r, "E", notes, fill=base_fill)

        ws.row_dimensions[r].height = ROW_HEIGHT_DATA

    return start_row + len(rows) - 1 if rows else start_row


def format_sheet(ws, last_row):
    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def build_tab(wb, tab_data, first_sheet=False):
    tab_name = tab_data.get("tab_name", "Storyboard")[:31]
    ws = wb.active if first_sheet else wb.create_sheet()
    ws.title = tab_name

    build_utility_bar(ws, tab_data.get("script_url", ""), tab_data.get("source_folder_url", ""))
    build_header(ws)
    last_row = build_rows(ws, tab_data.get("rows", []))
    format_sheet(ws, last_row)
    return ws


def build_legend(wb):
    ws = wb.create_sheet("Legend")
    set_cell(ws, 1, "A", "Visual Type", font=Font(bold=True, color="FFFFFF"),
             fill=PatternFill("solid", fgColor="1F2937"), align=WRAP_CENTER)
    set_cell(ws, 1, "B", "Color", font=Font(bold=True, color="FFFFFF"),
             fill=PatternFill("solid", fgColor="1F2937"), align=WRAP_CENTER)
    set_cell(ws, 1, "C", "Description", font=Font(bold=True, color="FFFFFF"),
             fill=PatternFill("solid", fgColor="1F2937"), align=WRAP_CENTER)

    descriptions = {
        "TALKING HEAD": "On-camera presenter",
        "B-ROLL": "Supporting footage / cutaway video",
        "CUSTOM GRAPHIC": "Motion graphic or designed visual",
        "SCREENCAST": "Screen recording of software/platform",
        "TESTIMONIAL": "Student/customer testimonial, video or card",
        "TALKING HEAD + LOWER THIRD": "Presenter with CTA text overlay",
        "B-ROLL / GRAPHIC": "Editor's choice between the two",
        "GRAPHIC / SCREENCAST": "Editor's choice between the two",
    }
    for i, (vtype, color) in enumerate(VISUAL_TYPES.items(), start=2):
        set_cell(ws, i, "A", vtype, font=Font(bold=True, color="FFFFFF"),
                 fill=PatternFill("solid", fgColor=color), align=WRAP)
        set_cell(ws, i, "B", f"#{color}", align=WRAP_CENTER)
        set_cell(ws, i, "C", descriptions.get(vtype, ""), align=WRAP)
        ws.row_dimensions[i].height = 24

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 45
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 build_storyboard.py input.json output.xlsx", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]
    with open(input_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    tabs = data.get("tabs", [])
    if not tabs:
        raise ValueError("Input JSON must contain at least one entry in 'tabs'.")

    for i, tab_data in enumerate(tabs):
        build_tab(wb, tab_data, first_sheet=(i == 0))

    build_legend(wb)
    wb.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
