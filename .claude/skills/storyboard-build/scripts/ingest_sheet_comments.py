#!/usr/bin/env python3
"""ingest_sheet_comments.py — pull reviewer comments out of a board .xlsx into storyboard.json.

The .xlsx rendered by sb_to_xlsx.py is the collaboration surface: people open it in Google
Sheets (or Excel) and comment on a row's cell. This reads those comment threads back out and
appends them to the matching row's `feedback[]`, so the board stays the source of truth and
`generate_row.py plan` lists them as open work.

    ingest_sheet_comments.py <storyboard.json> <workbook.xlsx> [--apply]

Without --apply it prints what it found and changes nothing.

Reads both comment formats:
  * Google Sheets export  — legacy `xl/comments*.xml`, where one cell's <text> holds a whole
    thread as "======\\nID#<id>\\n<author>  (<when>)\\n<body>" blocks, one per post.
  * Excel threaded comments — `xl/threadedComments/*.xml` + `xl/persons/person.xml`.
Plain old Excel notes (an <authors> list and free text) are read too, with no stable id.

A comment maps to a board row through the workbook's own "Line #" column, not through cell
position, so it survives the section header rows that sb_to_xlsx.py interleaves. Comments on a
cell with no line number (a section banner, the legend) are reported as unmapped and skipped.

Idempotent: a thread whose `comment_id` is already on the row is skipped, so re-running after
new comments only adds the new ones. Never changes a row's status — a comment is a request, and
whether it is addressed is decided by regenerating the row, not by reading the sheet.

Because sb_to_xlsx.py overwrites the workbook in place, --archive (on by default) first copies
the commented workbook into <project>/feedback/ so the threads survive the next rebuild.
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
import uuid
import xml.etree.ElementTree as ET
import zipfile

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
M_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
TC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"

LINE_COL_RE = re.compile(r"^\s*(line\s*#?|row\s*#?|#)\s*$", re.I)
# a thread opens with "======"; each REPLY inside it is introduced by "------"
BLOCK_SEP_RE = re.compile(r"^(={4,}|-{4,})\s*$")
ID_RE = re.compile(r"^ID#(\S+)\s*$")
AUTHOR_RE = re.compile(r"^(.*?)\s{2,}\((\d{4}-\d{2}-\d{2})[ T]([\d:]+)\)\s*$")


def die(msg):
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(2)


# ---------------------------------------------------------------- workbook plumbing

def _rels(zf, part):
    """Relationship id -> resolved zip path, for the given part (e.g. 'xl/worksheets/sheet1.xml')."""
    base, name = os.path.split(part)
    rels_path = f"{base}/_rels/{name}.rels"
    if rels_path not in zf.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(zf.read(rels_path)):
        target = rel.get("Target", "")
        if target.startswith("/"):
            resolved = target.lstrip("/")
        else:
            resolved = os.path.normpath(os.path.join(base, target)).replace(os.sep, "/")
        out[rel.get("Id")] = (rel.get("Type", ""), resolved)
    return out


def sheet_parts(zf):
    """[(sheet name, worksheet path, {'comments': path|None, 'threaded': [paths]})] in book order."""
    book = ET.fromstring(zf.read("xl/workbook.xml"))
    book_rels = _rels(zf, "xl/workbook.xml")
    out = []
    for sh in book.find(f"{{{M_NS}}}sheets"):
        rid = sh.get(f"{{{R_NS}}}id")
        if rid not in book_rels:
            continue
        ws_path = book_rels[rid][1]
        attached = {"comments": None, "threaded": []}
        for rel_type, target in _rels(zf, ws_path).values():
            if rel_type.endswith("/comments"):
                attached["comments"] = target
            elif rel_type.endswith("/threadedComment"):
                attached["threaded"].append(target)
        out.append((sh.get("name"), ws_path, attached))
    return out


def persons(zf):
    """Excel threaded-comment personId -> display name."""
    path = "xl/persons/person.xml"
    if path not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(path))
    return {p.get("id"): p.get("displayName") or "" for p in root}


# ---------------------------------------------------------------- comment parsing

def _to_local(stamp):
    """Sheets exports comment times in UTC; everything else on the board (registered_at, file
    mtimes) is local, and local is also what the commenter saw on screen. Convert once, here."""
    if not stamp or len(stamp) < 16:
        return stamp
    try:
        utc = dt.datetime.fromisoformat(stamp.replace("Z", "")).replace(tzinfo=dt.timezone.utc)
    except ValueError:
        return stamp
    return utc.astimezone().replace(tzinfo=None).isoformat(timespec="seconds")


def _post(author, when, body, when_at=None):
    """`when` is the date the note was written; `when_at` keeps the clock time too, so a rebuild
    the same day can still tell whether the art came before or after the note."""
    local = _to_local(when_at) if when_at else None
    return {"author": (author or "").strip(), "when": (local or when or "")[:10] or when,
            "when_at": local or when_at or when, "text": (body or "").strip()}


def parse_google_thread(text, fallback_author):
    """Split a Google-exported cell comment into its posts. Returns [] for empty text."""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    blocks, current = [], []
    for line in lines:
        if BLOCK_SEP_RE.match(line):
            if current:
                blocks.append(current)
            current = []
        else:
            current.append(line)
    if current:
        blocks.append(current)

    posts = []
    for block in blocks:
        cid, author, when, when_at, body = None, None, None, None, list(block)
        if body and ID_RE.match(body[0].strip()):
            cid = ID_RE.match(body[0].strip()).group(1)
            body = body[1:]
        if body:
            m = AUTHOR_RE.match(body[0].rstrip())
            if m:
                author, when, when_at = m.group(1), m.group(2), f"{m.group(2)}T{m.group(3)}"
                body = body[1:]
        post = _post(author or fallback_author, when, "\n".join(body), when_at)
        post["id"] = cid
        if post["text"] or cid:
            posts.append(post)
    if not posts and (text or "").strip():
        posts = [dict(_post(fallback_author, None, text), id=None)]
    return posts


def read_legacy_comments(zf, path, sheet_name):
    """Threads from xl/comments*.xml. Google puts a whole thread in one cell's text."""
    root = ET.fromstring(zf.read(path))
    authors = [a.text or "" for a in root.find(f"{{{M_NS}}}authors")] if root.find(f"{{{M_NS}}}authors") is not None else []
    threads = []
    clist = root.find(f"{{{M_NS}}}commentList")
    for c in (clist if clist is not None else []):
        ref = c.get("ref")
        try:
            fallback = authors[int(c.get("authorId") or 0)]
        except (IndexError, ValueError):
            fallback = ""
        text = "".join(t.text or "" for t in c.iter(f"{{{M_NS}}}t"))
        posts = parse_google_thread(text, fallback)
        if posts:
            threads.append({"sheet": sheet_name, "cell": ref, "posts": posts})
    return threads


def read_threaded_comments(zf, paths, people, sheet_name):
    """Threads from Excel's threadedComments parts (head post + replies via parentId)."""
    by_id, order = {}, []
    for path in paths:
        for tc in ET.fromstring(zf.read(path)):
            cid = (tc.get("id") or "").strip("{}")
            text_el = tc.find(f"{{{TC_NS}}}text")
            post = _post(people.get(tc.get("personId"), ""), (tc.get("dT") or "")[:10],
                         text_el.text if text_el is not None else "", (tc.get("dT") or "")[:19])
            post["id"] = cid
            rec = {"ref": tc.get("ref"), "parent": (tc.get("parentId") or "").strip("{}"), "post": post}
            by_id[cid] = rec
            order.append(cid)
    threads = {}
    for cid in order:
        rec = by_id[cid]
        head = rec["parent"] or cid
        threads.setdefault(head, {"sheet": sheet_name, "cell": by_id.get(head, rec)["ref"], "posts": []})
        threads[head]["posts"].append(rec["post"])
    return list(threads.values())


def extract_threads(xlsx_path, only_sheet=None):
    with zipfile.ZipFile(xlsx_path) as zf:
        people = persons(zf)
        found = []
        for name, _ws, attached in sheet_parts(zf):
            if only_sheet and name != only_sheet:
                continue
            if attached["threaded"]:
                found += read_threaded_comments(zf, attached["threaded"], people, name)
            elif attached["comments"]:
                found += read_legacy_comments(zf, attached["comments"], name)
    return found


# ---------------------------------------------------------------- row mapping

def line_number_index(ws):
    """1-based column index of the board's line-number column ('Line #'), else column 1."""
    for col in range(1, min(ws.max_column, 12) + 1):
        val = ws.cell(1, col).value
        if isinstance(val, str) and LINE_COL_RE.match(val):
            return col
    return 1


def map_rows(xlsx_path, threads, row_col=None):
    """Attach a board row number to each thread, via the workbook's own line-number column."""
    try:
        import openpyxl
    except ImportError:
        die("openpyxl is required: python3 -m pip install openpyxl")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    idx_cache = {}
    for th in threads:
        ws = wb[th["sheet"]]
        col = row_col or idx_cache.setdefault(th["sheet"], line_number_index(ws))
        try:
            cell_row = int(re.sub(r"\D", "", th["cell"]))
        except ValueError:
            th["row"] = None
            continue
        val = ws.cell(cell_row, col).value
        try:
            th["row"] = int(float(val))
        except (TypeError, ValueError):
            th["row"] = None
    return threads


# ---------------------------------------------------------------- board merge

def thread_text(posts):
    """Head post verbatim; replies appended as '↳ author: text' so nothing is lost."""
    if not posts:
        return ""
    out = [posts[0]["text"]]
    for reply in posts[1:]:
        who = reply["author"] or "reply"
        out.append(f"↳ {who}: {reply['text']}")
    return "\n".join(x for x in out if x.strip())


def already_ingested(row, comment_id):
    if not comment_id:
        return False
    return any((fb.get("comment_id") == comment_id) for fb in (row.get("feedback") or []))


def main():
    ap = argparse.ArgumentParser(description="Ingest .xlsx reviewer comments into storyboard.json")
    ap.add_argument("storyboard_json")
    ap.add_argument("workbook")
    ap.add_argument("--apply", action="store_true", help="write the feedback (default: dry run)")
    ap.add_argument("--who", help="override the commenter name recorded on every entry")
    ap.add_argument("--sheet", help="only read comments from this worksheet")
    ap.add_argument("--row-col", type=int, help="1-based column holding the board row number")
    ap.add_argument("--json", dest="json_out", help="also write the raw extraction to this path")
    ap.add_argument("--archive", metavar="DIR", default="feedback",
                    help="copy the commented workbook here first (default: <project>/feedback)")
    ap.add_argument("--no-archive", action="store_true", help="do not copy the workbook aside")
    args = ap.parse_args()

    if not os.path.exists(args.workbook):
        die(f"no such workbook: {args.workbook}")
    with open(args.storyboard_json) as f:
        doc = json.load(f)
    rows_by_n = {r["n"]: r for r in doc["rows"]}

    threads = map_rows(args.workbook, extract_threads(args.workbook, args.sheet), args.row_col)
    if not threads:
        print("no comments found in the workbook")
        return

    # Newest last, so a row's feedback reads in the order it was written.
    threads.sort(key=lambda t: (t["row"] or 0, t["posts"][0].get("when") or ""))

    project_dir = os.path.dirname(os.path.abspath(args.storyboard_json))
    archived = None
    if args.apply and not args.no_archive:
        dest_dir = args.archive if os.path.isabs(args.archive) else os.path.join(project_dir, args.archive)
        os.makedirs(dest_dir, exist_ok=True)
        stamp = threads[0]["posts"][0].get("when") or dt.date.today().isoformat()
        base = os.path.basename(args.workbook)
        archived = os.path.join(dest_dir, base if base.startswith(stamp) else f"{stamp}_{base}")
        if not os.path.exists(archived):
            shutil.copy2(args.workbook, archived)

    added, skipped, unmapped = [], [], []
    for th in threads:
        head = th["posts"][0]
        cid = head.get("id")
        text = thread_text(th["posts"])
        if not text:
            continue
        if th["row"] is None or th["row"] not in rows_by_n:
            unmapped.append(th)
            continue
        row = rows_by_n[th["row"]]
        if already_ingested(row, cid):
            skipped.append(th)
            continue
        entry = {
            "who": args.who or f"{head['author'] or 'reviewer'} (sheet)",
            "when": head.get("when") or dt.date.today().isoformat(),
            "when_at": head.get("when_at") or head.get("when") or dt.date.today().isoformat(),
            "text": text,
            "id": uuid.uuid4().hex[:8],
            "rows": [th["row"]],
            "source": "sheet-comment",
            "comment_id": cid,
            "cell": th["cell"],
        }
        th["entry"] = entry
        added.append(th)

    verb = "added" if args.apply else "would add"
    print(f"{args.workbook}: {len(threads)} thread(s) — {verb} {len(added)}, "
          f"already on board {len(skipped)}, unmapped {len(unmapped)}")
    for th in added:
        # the whole note on one line: a first line is often just the @mention
        flat = " / ".join(x.strip() for x in th["entry"]["text"].split("\n") if x.strip())
        print(f"  row {th['row']:>3}  {th['entry']['who']:<24} {th['cell']:<5} {flat[:150]}")
    for th in unmapped:
        print(f"  (skip) {th['cell']} on '{th['sheet']}' has no board row number: "
              f"{thread_text(th['posts'])[:60]}")

    if args.json_out:
        with open(args.json_out, "w") as f:
            json.dump({"workbook": os.path.abspath(args.workbook),
                       "extracted": dt.date.today().isoformat(),
                       "threads": [{k: v for k, v in th.items() if k != "entry"} for th in threads]},
                      f, indent=2, ensure_ascii=False)
        print(f"  extraction written to {args.json_out}")

    if not args.apply:
        print("\ndry run — re-run with --apply to write these to the board")
        return

    for th in added:
        rows_by_n[th["row"]].setdefault("feedback", []).append(th["entry"])
    with open(args.storyboard_json, "w") as f:
        json.dump(doc, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"\nwrote {len(added)} feedback entr{'y' if len(added) == 1 else 'ies'} to {args.storyboard_json}")
    if archived:
        print(f"commented workbook archived to {archived}")
    print("next: generate_row.py plan storyboard.json   (lists them as open feedback)")


if __name__ == "__main__":
    main()
