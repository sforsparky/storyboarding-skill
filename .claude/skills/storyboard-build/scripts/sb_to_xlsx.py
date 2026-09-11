#!/usr/bin/env python3
"""Render storyboard.json -> a formatted .xlsx — the board you open directly in Google Sheets.

Layout mirrors the working storyboard: section header rows, a row-number column tied to
storyboard.json ids, an Engine column (the /storyboard-motion route), and the thumbnail embedded
IN the Visual cell once a row has a local poster (the type label shows until then). No Status
column. This .xlsx is the collaboration surface — open it in Google Sheets, comment, and share.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, os.path.dirname(os.path.realpath(__file__)))
from lib_types import VISUAL_TYPES, TALKING_HEAD_TYPES, normalize_type, default_motion_engine

WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
COLS = {"A":"Line #","B":"Visual","C":"Script","D":"Visual Direction","E":"Notes","F":"Engine"}
WIDTHS = {"A":7,"B":28,"C":46,"D":40,"E":28,"F":14}
ROW_H = 96

def cell(ws,r,c,v,font=None,fill=None,align=WRAP):
    x=ws[f"{c}{r}"]; x.value=v; x.alignment=align
    if font:x.font=font
    if fill:x.fill=fill
    return x

def main():
    sb_path, out_path = sys.argv[1], sys.argv[2]
    doc = json.load(open(sb_path))
    proj_dir = os.path.dirname(os.path.abspath(sb_path))
    wb = Workbook(); ws = wb.active; ws.title = doc.get("project","Storyboard")[:31]

    for c,label in COLS.items():
        cell(ws,1,c,label,font=Font(bold=True,color="FFFFFF"),
             fill=PatternFill("solid",fgColor="1F2937"),align=CENTER)
    ws.row_dimensions[1].height=22
    for c,w in WIDTHS.items(): ws.column_dimensions[c].width=w
    ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False

    r=2; last_section=None; i=0
    for row in doc["rows"]:
        sec=row.get("section")
        if sec and sec!=last_section:
            ws.merge_cells(f"A{r}:F{r}")
            cell(ws,r,"A",f"▶  {sec}",font=Font(bold=True,color="FFFFFF",size=12),
                 fill=PatternFill("solid",fgColor="0F172A"),align=Alignment(vertical="center",horizontal="left"))
            ws.row_dimensions[r].height=26; last_section=sec; r+=1
        vtype=normalize_type(row["visual_type"]); color=VISUAL_TYPES[vtype]
        zebra=PatternFill("solid",fgColor="F2F2F2") if i%2 else None
        cell(ws,r,"A",row["n"],align=CENTER,fill=zebra)
        eng = row.get("motion_engine") or default_motion_engine(vtype)
        # find a poster (own asset or the reused row's)
        poster=None; src=row
        if row.get("reuse_of"):
            src=next((x for x in doc["rows"] if x["id"]==row["reuse_of"]), row)
        for a in src.get("assets",[]):
            if a.get("poster"): poster=os.path.join(proj_dir,a["poster"])
        # Visual cell: the thumbnail replaces the type label once a poster exists (type color stays as fill)
        label = "" if poster and os.path.exists(poster) else vtype + (f"\n↩ reuse {row['reuse_of']}" if row.get("reuse_of") else "")
        cell(ws,r,"B",label,font=Font(bold=True,color="FFFFFF"),
             fill=PatternFill("solid",fgColor=color),align=CENTER)
        if poster and os.path.exists(poster):
            try:
                img=XLImage(poster); img.height=ROW_H-8; img.width=int((ROW_H-8)*16/9)
                ws.add_image(img,f"B{r}")
            except Exception: pass
        cell(ws,r,"C",row.get("script",""),fill=zebra)
        cell(ws,r,"D",row.get("visual_direction","") if vtype not in TALKING_HEAD_TYPES else "",
             font=Font(color=color),fill=zebra)
        cell(ws,r,"E","" if vtype in TALKING_HEAD_TYPES else row.get("notes",""),fill=zebra)
        cell(ws,r,"F","" if eng=="none" else eng,align=CENTER,fill=zebra)
        ws.row_dimensions[r].height=ROW_H; r+=1; i+=1

    lg=wb.create_sheet("Legend")
    cell(lg,1,"A","Visual Type",font=Font(bold=True,color="FFFFFF"),fill=PatternFill("solid",fgColor="1F2937"),align=CENTER)
    cell(lg,1,"B","Color",font=Font(bold=True,color="FFFFFF"),fill=PatternFill("solid",fgColor="1F2937"),align=CENTER)
    for idx,(vt,cl) in enumerate(VISUAL_TYPES.items(),start=2):
        cell(lg,idx,"A",vt,font=Font(bold=True,color="FFFFFF"),fill=PatternFill("solid",fgColor=cl))
        cell(lg,idx,"B",f"#{cl}",align=CENTER)
    lg.column_dimensions["A"].width=28; lg.column_dimensions["B"].width=12
    lg.sheet_view.showGridLines=False
    wb.save(out_path); print("Saved",out_path)

if __name__=="__main__":
    main()
